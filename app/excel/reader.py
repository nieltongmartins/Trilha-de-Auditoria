"""Converte arquivos ``.xlsx`` em snapshots lógicos comparáveis.

O caminho rápido lê diretamente o XML interno do XLSX, evitando a criação de
objetos ``Cell`` do openpyxl. Para estruturas de fórmula não suportadas pelo
leitor direto, há fallback automático para o leitor openpyxl original.
"""

from __future__ import annotations

from io import BytesIO
import logging
from pathlib import Path, PurePosixPath
from typing import Callable, TypeAlias
import xml.etree.ElementTree as ET
import zipfile

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles.numbers import BUILTIN_FORMATS, is_date_format
from openpyxl.utils.datetime import (
    CALENDAR_MAC_1904,
    CALENDAR_WINDOWS_1900,
    from_ISO8601,
    from_excel,
)

from app.excel.formulas import normalize_formula_value


logger = logging.getLogger("auditoria_excel.reader")

CellValue: TypeAlias = str | int | float | bool | None
SheetSnapshot: TypeAlias = dict[str, CellValue]
Snapshot: TypeAlias = dict[str, SheetSnapshot]

_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PKG_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_NS = {"m": _MAIN_NS, "r": _REL_NS, "pr": _PKG_REL_NS}


class _FastReaderUnsupported(RuntimeError):
    """Sinaliza recurso raro em que o fallback openpyxl preserva melhor a semântica."""


def _tag(local: str) -> str:
    return f"{{{_MAIN_NS}}}{local}"


_FORMULA_TAG = _tag("f")
_VALUE_TAG = _tag("v")
_INLINE_STRING_TAG = _tag("is")


def _read_openpyxl(path: Path) -> Snapshot:
    """Fallback equivalente ao leitor original."""
    workbook = load_workbook(
        filename=path,
        read_only=True,
        data_only=False,
        keep_links=False,
    )
    try:
        snapshot: Snapshot = {}
        for worksheet in workbook.worksheets:
            cells: SheetSnapshot = {}
            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is not None:
                        cells[cell.coordinate] = normalize_formula_value(cell.value)
            snapshot[worksheet.title] = cells
        return snapshot
    finally:
        workbook.close()


def _cast_number(value: str) -> int | float:
    text = value.strip()
    if not text:
        raise ValueError("valor numérico vazio")
    if all(char not in text for char in ".eE"):
        try:
            return int(text)
        except ValueError:
            pass
    return float(text)


def _all_text(element: ET.Element | None) -> str:
    if element is None:
        return ""
    return "".join(node.text or "" for node in element.iter() if node.tag == _tag("t"))


def _shared_strings(archive: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in archive.namelist():
        return []
    result: list[str] = []
    with archive.open("xl/sharedStrings.xml") as stream:
        for event, element in ET.iterparse(stream, events=("end",)):
            if element.tag == _tag("si"):
                result.append(_all_text(element))
                element.clear()
    return result


def _workbook_epoch(archive: zipfile.ZipFile) -> object:
    root = ET.fromstring(archive.read("xl/workbook.xml"))
    workbook_pr = root.find("m:workbookPr", _NS)
    if workbook_pr is not None and workbook_pr.get("date1904") in {"1", "true", "True"}:
        return CALENDAR_MAC_1904
    return CALENDAR_WINDOWS_1900


def _date_styles(archive: zipfile.ZipFile) -> set[int]:
    if "xl/styles.xml" not in archive.namelist():
        return set()
    root = ET.fromstring(archive.read("xl/styles.xml"))
    custom_formats: dict[int, str] = {}
    num_fmts = root.find("m:numFmts", _NS)
    if num_fmts is not None:
        for num_fmt in num_fmts.findall("m:numFmt", _NS):
            try:
                num_fmt_id = int(num_fmt.get("numFmtId", ""))
            except ValueError:
                continue
            format_code = num_fmt.get("formatCode")
            if format_code is not None:
                custom_formats[num_fmt_id] = format_code

    result: set[int] = set()
    cell_xfs = root.find("m:cellXfs", _NS)
    if cell_xfs is None:
        return result
    for index, xf in enumerate(cell_xfs.findall("m:xf", _NS)):
        try:
            num_fmt_id = int(xf.get("numFmtId", "0"))
        except ValueError:
            num_fmt_id = 0
        format_code = custom_formats.get(num_fmt_id) or BUILTIN_FORMATS.get(num_fmt_id)
        if format_code and is_date_format(format_code):
            result.add(index)
    return result


def _sheet_targets(archive: zipfile.ZipFile) -> list[tuple[str, str]]:
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    rels = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rel_targets: dict[str, str] = {}
    for rel in rels.findall("pr:Relationship", _NS):
        rel_id = rel.get("Id")
        target = rel.get("Target")
        if rel_id and target:
            rel_targets[rel_id] = target

    result: list[tuple[str, str]] = []
    sheets = workbook.find("m:sheets", _NS)
    if sheets is None:
        return result
    for sheet in sheets.findall("m:sheet", _NS):
        title = sheet.get("name")
        rel_id = sheet.get(f"{{{_REL_NS}}}id")
        if not title or not rel_id or rel_id not in rel_targets:
            continue
        target = rel_targets[rel_id].lstrip("/")
        if not target.startswith("xl/"):
            target = str(PurePosixPath("xl") / target)
        result.append((title, target))
    return result


def _formula_element_value(
    formula: ET.Element | None,
    coordinate: str,
    shared_formulas: dict[str, tuple[str, str]],
) -> str | None:
    if formula is None:
        return None
    formula_type = formula.get("t")
    if formula_type in {"array", "dataTable"}:
        raise _FastReaderUnsupported(f"fórmula {formula_type} em {coordinate}")

    text = formula.text or ""
    if formula_type == "shared":
        shared_index = formula.get("si")
        if shared_index is None:
            raise _FastReaderUnsupported(f"fórmula compartilhada sem si em {coordinate}")
        if text:
            value = "=" + text
            shared_formulas[shared_index] = (coordinate, value)
            return value
        master = shared_formulas.get(shared_index)
        if master is None:
            raise _FastReaderUnsupported(
                f"fórmula compartilhada sem mestre conhecido em {coordinate}"
            )
        origin, master_formula = master
        return Translator(master_formula, origin=origin).translate_formula(coordinate)

    return "=" + text


def _formula_value(
    cell: ET.Element,
    coordinate: str,
    shared_formulas: dict[str, tuple[str, str]],
) -> str | None:
    """Compatibility helper used by the profiling baseline."""
    return _formula_element_value(cell.find("m:f", _NS), coordinate, shared_formulas)


def _convert_cell_value(
    cell: ET.Element,
    coordinate: str,
    formula_element: ET.Element | None,
    value_element: ET.Element | None,
    inline_element: ET.Element | None,
    shared_strings: list[str],
    date_styles: set[int],
    epoch: object,
    shared_formulas: dict[str, tuple[str, str]],
) -> object:
    formula = _formula_element_value(formula_element, coordinate, shared_formulas)
    if formula is not None:
        return formula

    cell_type = cell.get("t", "n")
    raw = value_element.text if value_element is not None else None

    if cell_type == "inlineStr":
        text = _all_text(inline_element)
        return text if text != "" else None
    if raw is None:
        return None
    if cell_type == "s":
        index = int(raw)
        return shared_strings[index]
    if cell_type == "b":
        return raw == "1"
    if cell_type in {"str", "e"}:
        return raw
    if cell_type == "d":
        return from_ISO8601(raw)

    value = _cast_number(raw)
    style_text = cell.get("s")
    if style_text is not None:
        try:
            style_id = int(style_text)
        except ValueError:
            style_id = -1
        if style_id in date_styles:
            return from_excel(value, epoch)
    return value


def _cell_value(
    cell: ET.Element,
    coordinate: str,
    shared_strings: list[str],
    date_styles: set[int],
    epoch: object,
    shared_formulas: dict[str, tuple[str, str]],
) -> object:
    """Converte uma célula após uma única visita a seus filhos XML diretos."""
    formula_element = value_element = inline_element = None
    for child in cell:
        # Element.find() selecionava a primeira ocorrência; mantenha a mesma
        # semântica até para XML incomum com filhos duplicados.
        if child.tag == _FORMULA_TAG and formula_element is None:
            formula_element = child
        elif child.tag == _VALUE_TAG and value_element is None:
            value_element = child
        elif child.tag == _INLINE_STRING_TAG and inline_element is None:
            inline_element = child
    return _convert_cell_value(
        cell,
        coordinate,
        formula_element,
        value_element,
        inline_element,
        shared_strings,
        date_styles,
        epoch,
        shared_formulas,
    )


def _cell_value_repeated_find(
    cell: ET.Element,
    coordinate: str,
    shared_strings: list[str],
    date_styles: set[int],
    epoch: object,
    shared_formulas: dict[str, tuple[str, str]],
) -> object:
    """Implementação anterior, mantida apenas como baseline de equivalência."""
    formula = _formula_value(cell, coordinate, shared_formulas)
    if formula is not None:
        return formula

    cell_type = cell.get("t", "n")
    value_element = cell.find("m:v", _NS)
    raw = value_element.text if value_element is not None else None
    if cell_type == "inlineStr":
        text = _all_text(cell.find("m:is", _NS))
        return text if text != "" else None
    if raw is None:
        return None
    if cell_type == "s":
        return shared_strings[int(raw)]
    if cell_type == "b":
        return raw == "1"
    if cell_type in {"str", "e"}:
        return raw
    if cell_type == "d":
        return from_ISO8601(raw)

    value = _cast_number(raw)
    style_text = cell.get("s")
    if style_text is not None:
        try:
            style_id = int(style_text)
        except ValueError:
            style_id = -1
        if style_id in date_styles:
            return from_excel(value, epoch)
    return value


def _read_fast_with(
    path: Path,
    cell_reader: Callable[
        [ET.Element, str, list[str], set[int], object, dict[str, tuple[str, str]]],
        object,
    ],
) -> Snapshot:
    with zipfile.ZipFile(path) as archive:
        shared = _shared_strings(archive)
        date_styles = _date_styles(archive)
        epoch = _workbook_epoch(archive)
        sheets = _sheet_targets(archive)
        snapshot: Snapshot = {}

        for title, target in sheets:
            cells: SheetSnapshot = {}
            shared_formulas: dict[str, tuple[str, str]] = {}
            with archive.open(target) as stream:
                for event, element in ET.iterparse(stream, events=("end",)):
                    if element.tag != _tag("c"):
                        continue
                    coordinate = element.get("r")
                    if coordinate:
                        value = cell_reader(
                            element,
                            coordinate,
                            shared,
                            date_styles,
                            epoch,
                            shared_formulas,
                        )
                        if value is not None:
                            cells[coordinate] = value  # type: ignore[assignment]
                    element.clear()
            snapshot[title] = cells
        return snapshot


def _read_fast(path: Path) -> Snapshot:
    return _read_fast_with(path, _cell_value)


def _read_fast_repeated_find(path: Path) -> Snapshot:
    """Executa o leitor anterior para testes e benchmarks, nunca em produção."""
    return _read_fast_with(path, _cell_value_repeated_find)


def read_workbook(path: str | Path) -> Snapshot:
    """Lê um ``.xlsx`` e preserva a semântica usada pela auditoria.

    O leitor XML é usado por padrão. Recursos raros cuja equivalência não pode
    ser garantida caem automaticamente no openpyxl original.
    """
    workbook_path = Path(path)
    if workbook_path.suffix.lower() != ".xlsx":
        raise ValueError("O leitor aceita somente arquivos .xlsx")

    try:
        snapshot = _read_fast(workbook_path)
        logger.debug("Leitura XLSX concluída modo=xml arquivo=%s", workbook_path.name)
        return snapshot
    except _FastReaderUnsupported as error:
        logger.info(
            "Leitor XML usou fallback openpyxl arquivo=%s motivo=%s",
            workbook_path.name,
            error,
        )
        return _read_openpyxl(workbook_path)
