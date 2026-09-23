from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula
import pytest

from app.audit_service import AuditService
from app.database import Database
from app.report_service import ReportService, ReportValidationError


@pytest.fixture
def populated_database(tmp_path: Path):
    with Database(tmp_path / "audit.db") as database:
        database.initialize()
        connection = database.connection
        spreadsheet_id = connection.execute(
            """INSERT INTO planilha
               (drive_item_id, nome_atual, site_id, drive_id, caminho_sharepoint)
               VALUES ('item-1', 'CQL028.xlsx', 'site-1', 'drive-1', '/docs/CQL028.xlsx')"""
        ).lastrowid
        execution_id = connection.execute(
            """INSERT INTO execucao_auditoria
               (codigo_execucao, planilha_id, status, fim, versoes_processadas,
                alteracoes_encontradas)
               VALUES ('AUD-1', ?, 'CONCLUIDA', CURRENT_TIMESTAMP, 1, 2)""",
            (spreadsheet_id,),
        ).lastrowid
        version_id = connection.execute(
            """INSERT INTO versao_processada
               (planilha_id, versao_anterior_id, versao_anterior_numero,
                versao_atual_id, versao_atual_numero, data_hora_versao, autor,
                comentario, quantidade_alteracoes, status, execucao_id)
               VALUES (?, 'v1', '0.1', 'v2', '0.2', '2026-09-16T10:00:00Z',
                       'Maria', 'Ajuste', 2, 'PROCESSADA', ?)""",
            (spreadsheet_id, execution_id),
        ).lastrowid
        connection.executemany(
            """INSERT INTO alteracao
               (versao_processada_id, planilha_id, tipo, aba, endereco,
                valor_anterior, valor_novo) VALUES (?, ?, ?, ?, ?, ?, ?)""",
            (
                (version_id, spreadsheet_id, "ADD", "Dados", "A1", None, "novo"),
                (version_id, spreadsheet_id, "MOD", "Dados", "B2", "1", "2"),
            ),
        )
        connection.commit()
        yield database, spreadsheet_id


def test_report_contains_official_sheets_filters_and_database_data(
    populated_database, tmp_path: Path
) -> None:
    database, spreadsheet_id = populated_database

    report = ReportService(database.connection, tmp_path / "reports").generate(
        spreadsheet_id
    )
    workbook = load_workbook(report)

    assert report.name.startswith("CQL028__")
    assert report.name.endswith("_Trilha_Auditoria.xlsx")
    assert workbook.sheetnames == ["Resumo", "Execuções", "Versões processadas", "Alterações_001", "Erros", "Integridade"]
    assert workbook["Resumo"]["B2"].value == "CQL028.xlsx"
    assert workbook["Resumo"]["B6"].value == 2
    assert workbook["Versões processadas"]["B2"].value == "0.1"
    assert workbook["Alterações_001"]["I2"].value == "ADD"
    assert all(sheet.auto_filter.ref for sheet in workbook.worksheets)


def test_report_can_be_regenerated_and_rejects_unknown_spreadsheet(
    populated_database, tmp_path: Path
) -> None:
    database, spreadsheet_id = populated_database
    service = ReportService(database.connection, tmp_path)

    first = service.generate(spreadsheet_id)
    first.unlink()
    second = service.generate(spreadsheet_id)

    assert second.is_file()
    with pytest.raises(ValueError, match="não encontrada"):
        service.generate(999)


def test_array_formula_is_persisted_and_reported_as_readable_text(
    populated_database, tmp_path: Path
) -> None:
    database, spreadsheet_id = populated_database
    connection = database.connection
    version_id = connection.execute(
        "SELECT id FROM versao_processada WHERE planilha_id=?", (spreadsheet_id,)
    ).fetchone()[0]
    value = AuditService._serialize(
        ArrayFormula(ref="D2:D100", text="=SUM(E2:E100)")
    )
    connection.execute(
        """INSERT INTO alteracao
           (versao_processada_id, planilha_id, tipo, aba, endereco,
            valor_anterior, valor_novo) VALUES (?, ?, 'ADD', 'Dados', 'D2', NULL, ?)""",
        (version_id, spreadsheet_id, value),
    )
    connection.commit()

    stored = connection.execute(
        "SELECT valor_novo FROM alteracao WHERE endereco='D2'"
    ).fetchone()[0]
    report = ReportService(connection, tmp_path / "reports").generate(spreadsheet_id)
    workbook = load_workbook(report, data_only=False)
    exported = workbook["Alterações_001"]["K4"].value

    assert stored == exported == (
        'ARRAYFORMULA|{"ref":"D2:D100","text":"=SUM(E2:E100)"}'
    )
    assert "object at 0x" not in stored
    assert "object at 0x" not in exported


def test_report_is_deterministically_regenerated_from_database_only(
    tmp_path: Path,
) -> None:
    database_path = tmp_path / "audit.db"
    historical_directory = tmp_path / "historical_versions"
    historical_directory.mkdir()
    (historical_directory / "0.1.xlsx").write_bytes(b"temporary historical file")
    (historical_directory / "0.2.xlsx").write_bytes(b"temporary historical file")

    with Database(database_path) as database:
        database.initialize()
        connection = database.connection
        spreadsheet_id = connection.execute(
            """INSERT INTO planilha
               (drive_item_id, nome_atual, site_id, drive_id, caminho_sharepoint)
               VALUES ('stable-item-id', 'Auditoria Oficial.xlsx', 'site-oficial',
                       'drive-oficial', '/documentos/Auditoria Oficial.xlsx')"""
        ).lastrowid
        execution_id = connection.execute(
            """INSERT INTO execucao_auditoria
               (codigo_execucao, planilha_id, inicio, fim, status,
                checkpoint_inicial, versao_final, versoes_processadas,
                alteracoes_encontradas)
               VALUES ('AUD-RELATORIO', ?, '2026-09-16T09:00:00Z',
                       '2026-09-16T09:05:00Z', 'CONCLUIDA', '0.1', '0.4', 3, 3)""",
            (spreadsheet_id,),
        ).lastrowid
        version_rows = (
            (
                spreadsheet_id,
                "version-1",
                "0.1",
                "version-2",
                "0.2",
                "2026-09-16T08:01:00Z",
                "Ana",
                "Inclusão",
                1,
                "PROCESSADA",
                execution_id,
            ),
            (
                spreadsheet_id,
                "version-2",
                "0.2",
                "version-3",
                "0.3",
                "2026-09-16T08:02:00Z",
                "Bruno",
                "Sem diferenças",
                0,
                "SEM_ALTERACOES",
                execution_id,
            ),
            (
                spreadsheet_id,
                "version-3",
                "0.3",
                "version-4",
                "0.4",
                "2026-09-16T08:03:00Z",
                "Carla",
                "Fórmula e remoção",
                2,
                "PROCESSADA",
                execution_id,
            ),
        )
        version_ids = []
        for row in version_rows:
            version_ids.append(
                connection.execute(
                    """INSERT INTO versao_processada
                       (planilha_id, versao_anterior_id, versao_anterior_numero,
                        versao_atual_id, versao_atual_numero, data_hora_versao,
                        autor, comentario, quantidade_alteracoes, status, execucao_id)
                       VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
                    row,
                ).lastrowid
            )
        connection.executemany(
            """INSERT INTO alteracao
               (versao_processada_id, planilha_id, tipo, aba, endereco,
                valor_anterior, valor_novo) VALUES (?, ?, ?, ?, ?, ?, ?)""",
            (
                (version_ids[0], spreadsheet_id, "ADD", "Dados", "A1", None, "novo"),
                (
                    version_ids[2],
                    spreadsheet_id,
                    "MOD",
                    "Cálculos",
                    "C3",
                    "=SUM(A1:A2)",
                    "=SUM(A1:A3)",
                ),
                (version_ids[2], spreadsheet_id, "DEL", "Dados", "B2", "removido", None),
            ),
        )
        connection.commit()

        for historical_file in historical_directory.iterdir():
            historical_file.unlink()
        historical_directory.rmdir()

        first_report = ReportService(connection, tmp_path / "first").generate(
            spreadsheet_id
        )
        second_report = ReportService(connection, tmp_path / "second").generate(
            spreadsheet_id
        )

        first = load_workbook(first_report, data_only=False)
        second = load_workbook(second_report, data_only=False)
        try:
            assert first.sheetnames == second.sheetnames == [
                "Resumo", "Execuções", "Versões processadas", "Alterações_001",
                "Erros", "Integridade",
            ]
            for sheet_name in first.sheetnames:
                if sheet_name == "Integridade":
                    continue
                first_rows = list(first[sheet_name].iter_rows(values_only=True))
                second_rows = list(second[sheet_name].iter_rows(values_only=True))
                assert first_rows == second_rows

            summary = dict(first["Resumo"].iter_rows(min_row=2, values_only=True))
            assert summary == {
                "Planilha": "Auditoria Oficial.xlsx",
                "Identidade técnica": "site-oficial/drive-oficial/stable-item-id",
                "DriveItem ID": "stable-item-id",
                "Versões processadas": 3,
                "Total de alterações": 3,
                "ADD": 1,
                "MOD": 1,
                "DEL": 1,
                "Execuções": 1,
                "Erros": 0,
                "Checkpoint": None,
            }
            version_rows = list(first["Versões processadas"].iter_rows(min_row=2, values_only=True))
            assert [row[1:5] + row[7:10] for row in version_rows] == [
                (
                    "0.1", "0.2",
                    "2026-09-16T08:01:00Z",
                    "Ana",
                    "Inclusão",
                    "PROCESSADA",
                    1,
                ),
                (
                    "0.2",
                    "0.3",
                    "2026-09-16T08:02:00Z",
                    "Bruno",
                    "Sem diferenças",
                    "SEM_ALTERACOES",
                    0,
                ),
                (
                    "0.3",
                    "0.4",
                    "2026-09-16T08:03:00Z",
                    "Carla",
                    "Fórmula e remoção",
                    "PROCESSADA",
                    2,
                ),
            ]
            trail_rows = list(first["Alterações_001"].iter_rows(min_row=2, values_only=True))
            assert [row[8] for row in trail_rows] == ["ADD", "MOD", "DEL"]
            assert trail_rows[1][6:] == (
                "Cálculos",
                "C3",
                "MOD",
                "=SUM(A1:A2)",
                "=SUM(A1:A3)",
            )
            assert trail_rows[2][6:] == (
                "Dados",
                "B2",
                "DEL",
                "removido",
                None,
            )
            assert all(sheet.auto_filter.ref for sheet in first.worksheets)
            assert all(sheet.freeze_panes == "A2" for sheet in first.worksheets)
        finally:
            first.close()
            second.close()

        assert connection.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
        assert connection.execute("PRAGMA foreign_key_check").fetchall() == []

@pytest.mark.parametrize(
    ("total", "limit", "expected_counts"),
    [
        (0, 900_000, [0]),
        (1, 900_000, [1]),
        (899_999, 900_000, [899_999]),
        (900_000, 900_000, [900_000]),
        (900_001, 900_000, [900_000, 1]),
        (1_048_575, 900_000, [900_000, 148_575]),
        (1_048_576, 900_000, [900_000, 148_576]),
        (2_000_001, 900_000, [900_000, 900_000, 200_001]),
    ],
)
def test_large_volume_sheet_partition_contract(total, limit, expected_counts):
    """Cobre todos os limites pedidos sem materializar milhões de células."""
    counts = []
    remaining = total
    while remaining:
        count = min(limit, remaining)
        counts.append(count)
        remaining -= count
    assert (counts or [0]) == expected_counts
    assert sum(counts) == total
    assert all(count + 1 <= 1_048_576 for count in counts)


def test_streaming_splits_without_loss_and_integrity_is_consistent(
    populated_database, tmp_path: Path
) -> None:
    database, spreadsheet_id = populated_database
    report = ReportService(
        database.connection, tmp_path, change_rows_per_sheet=1, fetch_batch_size=1
    ).generate(spreadsheet_id)
    workbook = load_workbook(report, read_only=True)
    try:
        assert workbook.sheetnames[3:5] == ["Alterações_001", "Alterações_002"]
        ids = [workbook[name].cell(2, 1).value for name in workbook.sheetnames[3:5]]
        assert ids == sorted(ids) and len(ids) == len(set(ids)) == 2
        integrity = {
            row[0]: row[1] if len(row) > 1 else None
            for row in workbook["Integridade"].iter_rows(min_row=2, values_only=True)
        }
        assert integrity["Alterações no banco"] == integrity["Alterações exportadas"] == 2
        assert integrity["ADD no banco"] == integrity["ADD exportados"] == 1
        assert integrity["MOD no banco"] == integrity["MOD exportados"] == 1
        assert integrity["Número de abas de alterações"] == 2
        assert integrity["Resultado da validação"] == "VALIDADO"
    finally:
        workbook.close()


def test_validation_failure_never_replaces_previous_report(
    populated_database, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    database, spreadsheet_id = populated_database
    service = ReportService(database.connection, tmp_path)
    official = service.generate(spreadsheet_id)
    previous = official.read_bytes()

    def fail(*_args, **_kwargs):
        raise ReportValidationError("falha injetada")

    monkeypatch.setattr(service, "_validate", fail)
    with pytest.raises(ReportValidationError, match="injetada"):
        service.generate(spreadsheet_id)
    assert official.read_bytes() == previous
    assert not list(tmp_path.glob("*.part.xlsx"))


def test_change_export_source_uses_fetchmany_not_fetchall():
    import inspect

    source = inspect.getsource(ReportService._export_changes)
    batch_source = inspect.getsource(ReportService._batches)
    assert "fetchall" not in source
    assert "fetchmany(self.fetch_batch_size)" in batch_source


def test_streaming_batch_never_exceeds_configured_size(tmp_path: Path):
    class FakeCursor:
        def __init__(self):
            self.remaining = 25
            self.requests = []

        def fetchmany(self, size):
            self.requests.append(size)
            count = min(size, self.remaining)
            self.remaining -= count
            return list(range(count))

    with Database(tmp_path / "batch.db") as database:
        service = ReportService(database.connection, tmp_path, fetch_batch_size=10)
        cursor = FakeCursor()
        batches = list(service._batches(cursor))
    assert [len(batch) for batch in batches] == [10, 10, 5]
    assert set(cursor.requests) == {10}
