from collections.abc import Mapping
from pathlib import Path

import pytest
from openpyxl import Workbook


def _save_workbook(
    path: Path,
    cells: Mapping[str, object],
    *,
    support_sheet: bool = True,
    support_value: str = "Original",
    new_sheet: bool = False,
) -> None:
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "Dados"
    for address, value in cells.items():
        data_sheet[address] = value

    if support_sheet:
        workbook.create_sheet("Apoio")["A1"] = support_value
    if new_sheet:
        workbook.create_sheet("Nova Aba")["B2"] = "Nova"

    workbook.save(path)
    workbook.close()


@pytest.fixture
def local_workbook_versions(tmp_path: Path) -> Path:
    """Gera versões Excel controladas sem manter binários no repositório."""
    versions = tmp_path / "CQL028"
    versions.mkdir()

    version_084 = {
        "A1": 10,
        "C3": "Pendente",
        "E5": False,
        "F6": "=SUM(A1:A10)",
    }
    version_085 = {
        "A1": 15,
        "B2": "OK",
        "D4": 0,
        "E5": False,
        "F6": "=SUM(A1:A20)",
    }
    version_086 = {
        "A1": 15,
        "B2": "OK",
        "D4": False,
        "F6": "=SUM(A1:A20)",
    }

    _save_workbook(versions / "0.84.xlsx", version_084)
    _save_workbook(
        versions / "0.85.xlsx", version_085, support_value="Atualizado"
    )
    _save_workbook(
        versions / "0.86.xlsx", version_086, support_value="Atualizado"
    )
    _save_workbook(
        versions / "0.87.xlsx",
        version_086,
        support_sheet=False,
        new_sheet=True,
    )
    return versions
